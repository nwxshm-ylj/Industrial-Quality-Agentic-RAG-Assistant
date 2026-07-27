from pathlib import Path
from tempfile import TemporaryDirectory

from app.multimodal.document_parser import parse_multimodal_document


class _MockOcrProvider:
    def __init__(self) -> None:
        self.calls = 0

    def extract_text(self, image_data_uri: str) -> str:
        assert image_data_uri.startswith("data:image/")
        self.calls += 1
        return "OCR wheel inspection result"


def _create_pdf(path: Path) -> None:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Native wheel inspection text")
    document.save(path)
    document.close()


def _create_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    text_box.text = "Torque station diagnostic flow"
    presentation.save(path)


def main() -> None:
    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)
        pdf_path = root / "manual.pdf"
        pptx_path = root / "training.pptx"
        try:
            _create_pdf(pdf_path)
            _create_pptx(pptx_path)
        except ModuleNotFoundError as exc:
            print(
                "SKIPPED multimodal parser integration test: "
                f"optional runtime dependency is unavailable ({exc.name})"
            )
            return

        ocr = _MockOcrProvider()
        pdf = parse_multimodal_document(
            pdf_path,
            doc_id="pdf-doc",
            source="manual.pdf",
            doc_type="manual",
            version="v1",
            assets_root=root / "assets",
            ocr_provider=ocr,
        )
        assert len(pdf["assets"]) == 1
        assert "Native wheel inspection text" in pdf["content"]
        assert "OCR wheel inspection result" in pdf["content"]
        assert Path(pdf["assets"][0]["metadata"]["asset_path"]).exists()
        assert pdf["assets"][0]["embedding_input"].images
        assert ocr.calls == 1

        pptx = parse_multimodal_document(
            pptx_path,
            doc_id="pptx-doc",
            source="training.pptx",
            doc_type="training",
            version="v2",
            assets_root=root / "assets",
        )
        assert len(pptx["assets"]) == 1
        assert pptx["assets"][0]["metadata"]["modality"] == "text"
        assert "Torque station diagnostic flow" in pptx["content"]

    print("Multimodal PDF/PPTX parser tests passed with mock OCR")


if __name__ == "__main__":
    main()
