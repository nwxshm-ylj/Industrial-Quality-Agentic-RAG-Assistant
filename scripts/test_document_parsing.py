from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory

from app.rag.loader import load_single_document
from app.rag.splitter import split_docs


def _create_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("轮毂识别异常", level=1)
    document.add_paragraph("先确认相机曝光、镜头污染和标定状态。")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "检查项"
    table.cell(0, 1).text = "判定标准"
    table.cell(1, 0).text = "曝光"
    table.cell(1, 1).text = "灰度均值 120-180"
    table.cell(2, 0).text = "标定"
    table.cell(2, 1).text = "偏差小于 0.2 mm"
    document.save(path)


def _create_pdf(path: Path) -> None:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Wheel inspection manual")
    page.insert_text((72, 100), "Check exposure and calibration before replacing hardware.")
    document.save(path)
    document.close()


def _create_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "扭矩工位诊断"
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(6), Inches(1))
    text_box.text = "先检查拧紧程序版本，再检查传感器零点。"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(3), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "信号"
    table.cell(0, 1).text = "范围"
    table.cell(1, 0).text = "扭矩"
    table.cell(1, 1).text = "100-120 Nm"
    table.cell(2, 0).text = "角度"
    table.cell(2, 1).text = "30-45 deg"
    presentation.save(path)


def _assert_structured_document(document: dict, extension: str) -> None:
    assert document["file_ext"] == extension
    assert document["parser"] == "structured-v1"
    assert document["content"].strip()
    assert document["sections"]
    assert all(section["text"].strip() for section in document["sections"])


def main() -> None:
    with TemporaryDirectory() as temp_dir:
        root = Path(temp_dir)

        markdown_path = root / "wheel_sop.md"
        rows = "\n".join(
            f"| 检查项{i} | 判定标准{i} |" for i in range(1, 18)
        )
        markdown_path.write_text(
            "# 轮毂识别异常\n\n"
            "优先检查成像链路，再检查识别规则。\n\n"
            "## 检查标准\n\n"
            "| 检查项 | 判定标准 |\n"
            "| --- | --- |\n"
            f"{rows}\n",
            encoding="utf-8",
        )
        markdown = load_single_document(str(markdown_path))
        _assert_structured_document(markdown, ".md")
        assert any(section["section_type"] == "heading" for section in markdown["sections"])
        assert any(section["section_type"] == "table" for section in markdown["sections"])

        markdown_chunks = split_docs([markdown], chunk_size=240, chunk_overlap=40)
        table_chunks = [
            chunk for chunk in markdown_chunks
            if chunk["metadata"]["section_type"] == "table"
        ]
        assert len(table_chunks) > 1, "长表格应按行拆分为多个 chunk"
        assert all("| 检查项 | 判定标准 |" in chunk["text"] for chunk in table_chunks)
        assert all("[章节] 轮毂识别异常 > 检查标准" in chunk["text"] for chunk in table_chunks)

        txt_path = root / "quality_rule.txt"
        txt_path.write_bytes("第一章 检查规则\n\n轮毂跳动量不得超过 0.2 mm。".encode("gb18030"))
        text_document = load_single_document(str(txt_path))
        _assert_structured_document(text_document, ".txt")
        assert "轮毂跳动量" in text_document["content"]

        docx_path = root / "inspection.docx"
        _create_docx(docx_path)
        docx_document = load_single_document(str(docx_path))
        _assert_structured_document(docx_document, ".docx")
        assert any(section["section_type"] == "table" for section in docx_document["sections"])
        assert "[章节] 轮毂识别异常" in "\n".join(
            chunk["text"] for chunk in split_docs([docx_document])
        )

        pdf_path = root / "manual.pdf"
        _create_pdf(pdf_path)
        pdf_document = load_single_document(str(pdf_path))
        _assert_structured_document(pdf_document, ".pdf")
        assert all(section["page_number"] == 1 for section in pdf_document["sections"])
        assert "[页码] 1" in split_docs([pdf_document])[0]["text"]

        pptx_path = root / "training.pptx"
        _create_pptx(pptx_path)
        pptx_document = load_single_document(str(pptx_path))
        _assert_structured_document(pptx_document, ".pptx")
        assert any(section["section_type"] == "table" for section in pptx_document["sections"])
        assert any(section["page_number"] == 1 for section in pptx_document["sections"])

        legacy_chunks = split_docs(
            [{"source": "legacy.md", "content": "旧入库流程。" * 180}]
        )
        assert legacy_chunks
        assert all(
            chunk["metadata"]["section_type"] == "legacy_text"
            for chunk in legacy_chunks
        )

    print("Structured parsing and chunking tests passed")


if __name__ == "__main__":
    main()
