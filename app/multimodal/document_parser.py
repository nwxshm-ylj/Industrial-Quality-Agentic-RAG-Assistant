from __future__ import annotations

from base64 import b64encode
from pathlib import Path

from app.multimodal.ocr import DocumentOcrProvider
from app.multimodal.types import MultimodalEmbeddingInput


MULTIMODAL_DOCUMENT_EXTENSIONS = {".pdf", ".pptx"}
_MAX_IMAGE_BYTES = 10 * 1024 * 1024


def parse_multimodal_document(
    file_path: str | Path,
    *,
    doc_id: str,
    source: str,
    doc_type: str,
    version: str,
    assets_root: str | Path,
    ocr_provider: DocumentOcrProvider | None = None,
    asset_id_namespace: str | None = None,
) -> dict:
    path = Path(file_path)
    extension = path.suffix.lower()
    if extension not in MULTIMODAL_DOCUMENT_EXTENSIONS:
        raise ValueError(f"multimodal parsing does not support: {extension}")
    operation_folder = asset_id_namespace or "initial"
    asset_dir = Path(assets_root) / doc_id / operation_folder
    asset_dir.mkdir(parents=True, exist_ok=True)
    try:
        assets = (
            _parse_pdf(
                path,
                doc_id=doc_id,
                source=source,
                doc_type=doc_type,
                version=version,
                asset_dir=asset_dir,
                ocr_provider=ocr_provider,
                asset_id_namespace=asset_id_namespace,
            )
            if extension == ".pdf"
            else _parse_pptx(
                path,
                doc_id=doc_id,
                source=source,
                doc_type=doc_type,
                version=version,
                asset_dir=asset_dir,
                ocr_provider=ocr_provider,
                asset_id_namespace=asset_id_namespace,
            )
        )
    except Exception:
        _remove_asset_dir(asset_dir)
        raise
    content = "\n\n".join(
        asset["text"].strip() for asset in assets if asset["text"].strip()
    ).strip()
    if not assets:
        _remove_asset_dir(asset_dir)
        raise ValueError(f"document contains no multimodal assets: {path.name}")
    if not content:
        content = f"Multimodal assets extracted from {source}"
    return {
        "source": source,
        "content": content,
        "file_ext": extension,
        "assets": assets,
        "asset_dir": str(asset_dir),
    }


def _parse_pdf(
    path: Path,
    **context,
) -> list[dict]:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    assets = []
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document):
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
            image_bytes = pixmap.tobytes("png")
            assets.append(
                _build_asset(
                    image_bytes=image_bytes,
                    mime_type="image/png",
                    native_text=page.get_text().strip(),
                    ordinal=page_index + 1,
                    asset_kind="page",
                    **context,
                )
            )
    return assets


def _parse_pptx(
    path: Path,
    **context,
) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(path)
    assets = []
    for slide_index, slide in enumerate(presentation.slides):
        text_parts = []
        image_values: list[tuple[bytes, str]] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    text_parts.append(value)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_values.append(
                    (shape.image.blob, shape.image.content_type or "image/png")
                )
        if image_values:
            for image_index, (image_bytes, mime_type) in enumerate(
                image_values[:5],
                start=1,
            ):
                assets.append(
                    _build_asset(
                        image_bytes=image_bytes,
                        mime_type=mime_type,
                        native_text="\n".join(text_parts),
                        ordinal=slide_index + 1,
                        asset_kind=f"slide_{image_index}",
                        **context,
                    )
                )
        elif text_parts:
            assets.append(
                _build_text_asset(
                    text="\n".join(text_parts),
                    ordinal=slide_index + 1,
                    asset_kind="slide",
                    **context,
                )
            )
    return assets


def _build_asset(
    *,
    image_bytes: bytes,
    mime_type: str,
    native_text: str,
    ordinal: int,
    asset_kind: str,
    doc_id: str,
    source: str,
    doc_type: str,
    version: str,
    asset_dir: Path,
    ocr_provider: DocumentOcrProvider | None,
    asset_id_namespace: str | None,
) -> dict:
    if len(image_bytes) > _MAX_IMAGE_BYTES:
        raise ValueError(
            f"rendered image exceeds 10 MB limit: {source} page/slide {ordinal}"
        )
    extension = _extension_for_mime(mime_type)
    namespace = f"_{asset_id_namespace}" if asset_id_namespace else ""
    asset_id = f"{doc_id}{namespace}_{asset_kind}_{ordinal:04d}"
    asset_path = asset_dir / f"{asset_id}{extension}"
    asset_path.write_bytes(image_bytes)
    image_data_uri = _to_data_uri(image_bytes, mime_type)
    ocr_text = ocr_provider.extract_text(image_data_uri) if ocr_provider else ""
    combined_text = _merge_text(native_text, ocr_text)
    return {
        "text": combined_text,
        "embedding_input": MultimodalEmbeddingInput(
            text=combined_text or None,
            images=(image_data_uri,),
        ),
        "metadata": {
            "doc_id": doc_id,
            "asset_id": asset_id,
            "page_number": ordinal,
            "modality": "text+image" if combined_text else "image",
            "source": source,
            "doc_type": doc_type,
            "version": version,
            "asset_path": str(asset_path),
            "mime_type": mime_type,
        },
    }


def _build_text_asset(
    *,
    text: str,
    ordinal: int,
    asset_kind: str,
    doc_id: str,
    source: str,
    doc_type: str,
    version: str,
    asset_dir: Path,
    ocr_provider: DocumentOcrProvider | None,
    asset_id_namespace: str | None,
) -> dict:
    del asset_dir, ocr_provider
    namespace = f"_{asset_id_namespace}" if asset_id_namespace else ""
    asset_id = f"{doc_id}{namespace}_{asset_kind}_{ordinal:04d}"
    return {
        "text": text,
        "embedding_input": MultimodalEmbeddingInput(text=text),
        "metadata": {
            "doc_id": doc_id,
            "asset_id": asset_id,
            "page_number": ordinal,
            "modality": "text",
            "source": source,
            "doc_type": doc_type,
            "version": version,
            "asset_path": None,
            "mime_type": "text/plain",
        },
    }


def _merge_text(native_text: str, ocr_text: str) -> str:
    native_text = native_text.strip()
    ocr_text = ocr_text.strip()
    if not native_text:
        return ocr_text
    if not ocr_text or ocr_text in native_text:
        return native_text
    return f"{native_text}\n\n[OCR]\n{ocr_text}"


def _to_data_uri(image_bytes: bytes, mime_type: str) -> str:
    return f"data:{mime_type};base64,{b64encode(image_bytes).decode('ascii')}"


def _extension_for_mime(mime_type: str) -> str:
    return {
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/png": ".png",
        "image/webp": ".webp",
        "image/bmp": ".bmp",
        "image/tiff": ".tiff",
    }.get(mime_type.lower(), ".bin")


def _remove_asset_dir(asset_dir: Path) -> None:
    if not asset_dir.exists():
        return
    for path in asset_dir.iterdir():
        if path.is_file():
            path.unlink()
    try:
        asset_dir.rmdir()
    except OSError:
        pass
