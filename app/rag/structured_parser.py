from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable


_MARKDOWN_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_MARKDOWN_TABLE_SEPARATOR = re.compile(
    r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$"
)
_NUMBERED_HEADING = re.compile(
    r"^(?:"
    r"第[一二三四五六七八九十百千万0-9]+[编章节部分条]|"
    r"[一二三四五六七八九十]+[、.]|"
    r"\d+(?:\.\d+){0,3}[、.\s]"
    r")"
)


def parse_text_document(path: Path) -> dict:
    text = _read_text_with_fallback(path)
    sections = _parse_markdown(text) if path.suffix.lower() == ".md" else _parse_plain_text(text)
    return _build_result(path, sections)


def parse_pdf_document(path: Path) -> dict:
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf

    sections: list[dict[str, Any]] = []
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            table_sections, table_boxes = _extract_pdf_tables(page, page_index)
            for block in page.get_text("blocks", sort=True):
                if len(block) < 5:
                    continue
                bbox = tuple(float(value) for value in block[:4])
                text = _normalize_text(str(block[4]))
                if not text or _inside_any_box(bbox, table_boxes):
                    continue
                sections.append(
                    _section(
                        text,
                        section_type="paragraph",
                        page_number=page_index,
                    )
                )
            sections.extend(table_sections)
    return _build_result(path, sections)


def parse_docx_document(path: Path) -> dict:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    sections: list[dict[str, Any]] = []
    heading_path: list[str] = []
    table_index = 0

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = _normalize_text(paragraph.text)
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            heading_level = _heading_level(style_name, text)
            if heading_level:
                heading_path = _update_heading_path(heading_path, heading_level, text)
                sections.append(
                    _section(
                        text,
                        section_type="heading",
                        heading_path=heading_path,
                        heading_level=heading_level,
                    )
                )
            else:
                sections.append(
                    _section(
                        text,
                        section_type="paragraph",
                        heading_path=heading_path,
                    )
                )
        elif child.tag.endswith("}tbl"):
            table_index += 1
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            table_text = _table_to_markdown(rows)
            if table_text:
                sections.append(
                    _section(
                        table_text,
                        section_type="table",
                        heading_path=heading_path,
                        table_index=table_index,
                    )
                )

    return _build_result(path, sections)


def parse_pptx_document(path: Path) -> dict:
    from pptx import Presentation

    presentation = Presentation(path)
    sections: list[dict[str, Any]] = []
    table_index = 0

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_title = _normalize_text(slide.shapes.title.text) if slide.shapes.title else ""
        heading_path = [slide_title] if slide_title else []
        if slide_title:
            sections.append(
                _section(
                    slide_title,
                    section_type="heading",
                    heading_path=heading_path,
                    heading_level=1,
                    page_number=slide_number,
                )
            )

        for shape in sorted(
            slide.shapes,
            key=lambda item: (int(getattr(item, "top", 0)), int(getattr(item, "left", 0))),
        ):
            if slide.shapes.title is not None and shape == slide.shapes.title:
                continue
            if getattr(shape, "has_table", False):
                table_index += 1
                rows = [
                    [shape.table.cell(row, col).text for col in range(len(shape.table.columns))]
                    for row in range(len(shape.table.rows))
                ]
                table_text = _table_to_markdown(rows)
                if table_text:
                    sections.append(
                        _section(
                            table_text,
                            section_type="table",
                            heading_path=heading_path,
                            page_number=slide_number,
                            table_index=table_index,
                        )
                    )
                continue

            text = _extract_pptx_shape_text(shape)
            if text:
                sections.append(
                    _section(
                        text,
                        section_type="paragraph",
                        heading_path=heading_path,
                        page_number=slide_number,
                    )
                )

    return _build_result(path, sections)


def _parse_markdown(text: str) -> list[dict[str, Any]]:
    lines = text.splitlines()
    sections: list[dict[str, Any]] = []
    heading_path: list[str] = []
    paragraph: list[str] = []
    index = 0
    table_index = 0

    def flush_paragraph() -> None:
        if not paragraph:
            return
        value = _normalize_text("\n".join(paragraph))
        paragraph.clear()
        if value:
            sections.append(
                _section(
                    value,
                    section_type="paragraph",
                    heading_path=heading_path,
                )
            )

    while index < len(lines):
        line = lines[index]
        heading = _MARKDOWN_HEADING.match(line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_path = _update_heading_path(heading_path, level, title)
            sections.append(
                _section(
                    title,
                    section_type="heading",
                    heading_path=heading_path,
                    heading_level=level,
                )
            )
            index += 1
            continue

        if _is_markdown_table_start(lines, index):
            flush_paragraph()
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                table_lines.append(lines[index])
                index += 1
            table_index += 1
            sections.append(
                _section(
                    "\n".join(table_lines),
                    section_type="table",
                    heading_path=heading_path,
                    table_index=table_index,
                )
            )
            continue

        if line.strip():
            paragraph.append(line.rstrip())
        else:
            flush_paragraph()
        index += 1

    flush_paragraph()
    return sections


def _parse_plain_text(text: str) -> list[dict[str, Any]]:
    sections: list[dict[str, Any]] = []
    heading_path: list[str] = []
    for block in re.split(r"\n\s*\n", text):
        value = _normalize_text(block)
        if not value:
            continue
        first_line = value.splitlines()[0].strip()
        if len(value.splitlines()) == 1 and _looks_like_heading(first_line):
            level = _numbered_heading_level(first_line)
            heading_path = _update_heading_path(heading_path, level, first_line)
            sections.append(
                _section(
                    first_line,
                    section_type="heading",
                    heading_path=heading_path,
                    heading_level=level,
                )
            )
        else:
            sections.append(
                _section(
                    value,
                    section_type="paragraph",
                    heading_path=heading_path,
                )
            )
    return sections


def _extract_pdf_tables(page: Any, page_number: int) -> tuple[list[dict], list[tuple[float, ...]]]:
    if not hasattr(page, "find_tables"):
        return [], []
    try:
        finder = page.find_tables()
    except Exception:
        return [], []

    sections = []
    boxes = []
    for table_index, table in enumerate(getattr(finder, "tables", []), start=1):
        rows = table.extract()
        text = _table_to_markdown(rows)
        if not text:
            continue
        bbox = tuple(float(value) for value in table.bbox)
        boxes.append(bbox)
        sections.append(
            _section(
                text,
                section_type="table",
                page_number=page_number,
                table_index=table_index,
            )
        )
    return sections, boxes


def _extract_pptx_shape_text(shape: Any) -> str:
    if getattr(shape, "shape_type", None) == 6 and hasattr(shape, "shapes"):
        parts = [
            _extract_pptx_shape_text(child)
            for child in sorted(
                shape.shapes,
                key=lambda item: (
                    int(getattr(item, "top", 0)),
                    int(getattr(item, "left", 0)),
                ),
            )
        ]
        return _normalize_text("\n".join(part for part in parts if part))
    if getattr(shape, "has_text_frame", False):
        return _normalize_text(shape.text)
    return ""


def _table_to_markdown(rows: Iterable[Iterable[Any]]) -> str:
    normalized_rows = []
    for row in rows:
        values = [_normalize_cell(value) for value in row]
        if any(values):
            normalized_rows.append(values)
    if not normalized_rows:
        return ""

    width = max(len(row) for row in normalized_rows)
    normalized_rows = [row + [""] * (width - len(row)) for row in normalized_rows]
    header = normalized_rows[0]
    body = normalized_rows[1:]
    result = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    result.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(result)


def _build_result(path: Path, sections: list[dict[str, Any]]) -> dict:
    sections = [section for section in sections if section.get("text", "").strip()]
    content = "\n\n".join(section["text"] for section in sections).strip()
    if not content:
        raise ValueError(f"文档内容为空: {path.name}")
    return {
        "source": path.name,
        "content": content,
        "file_ext": path.suffix.lower(),
        "sections": sections,
        "parser": "structured-v1",
    }


def _section(
    text: str,
    *,
    section_type: str,
    heading_path: Iterable[str] | None = None,
    heading_level: int | None = None,
    page_number: int | None = None,
    table_index: int | None = None,
) -> dict[str, Any]:
    return {
        "text": _normalize_text(text),
        "section_type": section_type,
        "heading_path": list(heading_path or []),
        "heading_level": heading_level,
        "page_number": page_number,
        "table_index": table_index,
    }


def _read_text_with_fallback(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别文本编码: {path.name}（支持 UTF-8/GB18030）")


def _normalize_text(value: str) -> str:
    value = value.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    return _normalize_text(str(value)).replace("|", "\\|").replace("\n", "<br>")


def _is_markdown_table_start(lines: list[str], index: int) -> bool:
    return (
        index + 1 < len(lines)
        and "|" in lines[index]
        and bool(_MARKDOWN_TABLE_SEPARATOR.match(lines[index + 1]))
    )


def _inside_any_box(block: tuple[float, ...], boxes: list[tuple[float, ...]]) -> bool:
    center_x = (block[0] + block[2]) / 2
    center_y = (block[1] + block[3]) / 2
    return any(
        box[0] <= center_x <= box[2] and box[1] <= center_y <= box[3]
        for box in boxes
    )


def _heading_level(style_name: str, text: str) -> int | None:
    match = re.search(r"(?:Heading|标题)\s*([1-6])", style_name, re.IGNORECASE)
    if match:
        return int(match.group(1))
    if _looks_like_heading(text):
        return _numbered_heading_level(text)
    return None


def _looks_like_heading(text: str) -> bool:
    return len(text) <= 80 and bool(_NUMBERED_HEADING.match(text.strip()))


def _numbered_heading_level(text: str) -> int:
    dotted = re.match(r"^(\d+(?:\.\d+)*)", text)
    if dotted:
        return min(dotted.group(1).count(".") + 1, 6)
    if re.match(r"^第.+[编章部分]", text):
        return 1
    if re.match(r"^第.+节", text):
        return 2
    return 3


def _update_heading_path(current: list[str], level: int, title: str) -> list[str]:
    level = max(1, min(level, 6))
    result = list(current[: level - 1])
    while len(result) < level - 1:
        result.append("")
    result.append(title)
    return [value for value in result if value]
